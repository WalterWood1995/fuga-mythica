# -*- coding: utf-8 -*-
"""Adds sheet「1」: the simplified per-word breakdown, with the root meaning
written once per family in a merged cell (the layout the user sketched by hand)."""
import io, sys, json, subprocess, openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
path = sys.argv[1]
rows = json.loads(subprocess.check_output(["node", "-e",
    "const M=require('./tools/de_morph.js');process.stdout.write(JSON.stringify(M.rows))"]).decode("utf-8"))

sys.path.insert(0, __import__("os").path.dirname(__file__))
from de_root_groups import cluster
fams = {}
for r in rows: fams.setdefault(r["famKey"], {"key": r["famKey"], "zh": r["rootZh"], "en": r["rootEn"]})
groups = cluster(list(fams.values()))
gid, glabel = {}, {}
order_keys = []
for g in groups:
    g.sort(key=lambda k: k.lower().lstrip("-"))
for n, g in enumerate(sorted(groups, key=lambda g: g[0].lower().lstrip("-"))):
    for k in g: gid[k] = n
    glabel[n] = " · ".join(g) if len(g) > 1 else ""
    order_keys += g
pos = {k: i for i, k in enumerate(order_keys)}
rows.sort(key=lambda r: (pos[r["famKey"]], r["word"].lower()))

wb = openpyxl.load_workbook(path)
if "1" in wb.sheetnames: del wb["1"]
ws = wb.create_sheet("1", 0)
H = ["前缀", "复合前件", "连接成分", "核心词形", "复合后件", "余部", "后缀", "完整单词", "冠词", "中文", "English", "词根义中文", "词根义英语", "形近义近的词根组"]
ws.append(H)
gold = PatternFill("solid", fgColor="E8B64C"); alt = PatternFill("solid", fgColor="F6EFE3")
top = Border(top=Side(style="thin", color="B8A07A"))
for c in ws[1]: c.font = Font(bold=True); c.fill = gold
# one merged meaning per family; if two vocab roots land in the same family, list both meanings
means = {}
for r in rows:
    m = means.setdefault(r["famKey"], ([], []))
    if r["rootZh"] not in m[0]: m[0].append(r["rootZh"])
    if r["rootEn"] not in m[1]: m[1].append(r["rootEn"])
r0, fam, shade = 2, None, False
for i, r in enumerate(rows):
    ws.append([r["pre"], r["comp1"], r["link"], r["core"], r["comp2"], r["tail"], r["suf"],
               r["word"], r["art"], r["zh"], r["en"], " / ".join(means[r["famKey"]][0]), " / ".join(means[r["famKey"]][1]), glabel[gid[r["famKey"]]]])
    row = i + 2
    if r["famKey"] != fam:                       # a new family starts: close the previous block
        if fam is not None and row - 1 > r0:
            ws.merge_cells(start_row=r0, start_column=12, end_row=row - 1, end_column=12)
            ws.merge_cells(start_row=r0, start_column=13, end_row=row - 1, end_column=13)
        fam, r0, shade = r["famKey"], row, not shade
        for c in ws[row]: c.border = top
    if shade:
        for c in ws[row]: c.fill = alt
last = len(rows) + 1
thick = Side(style="medium", color="8A6D33")
g0, cur = 2, None
for i, r in enumerate(rows + [None]):
    row = i + 2
    g = gid[r["famKey"]] if r else None
    if g != cur:
        if cur is not None and glabel[cur] and row - 1 > g0:
            ws.merge_cells(start_row=g0, start_column=14, end_row=row - 1, end_column=14)
        if r:
            for c in ws[row]: c.border = Border(top=thick)
        g0, cur = row, g
if last > r0:
    ws.merge_cells(start_row=r0, start_column=12, end_row=last, end_column=12)
    ws.merge_cells(start_row=r0, start_column=13, end_row=last, end_column=13)
for row in ws.iter_rows(min_row=2, min_col=12, max_col=14):
    for c in row: c.alignment = Alignment(vertical="center", wrap_text=True)
for col, w in zip("ABCDEFGHIJKLMN", [9, 11, 12, 11, 11, 6, 13, 22, 5, 20, 22, 16, 18, 22]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
wb.save(path)
nf = len({r["famKey"] for r in rows})
print(f"sheet 1: {len(rows)} 行, {nf} 个词根家族, {sum(1 for v in glabel.values() if v)} 个形近义近组")
