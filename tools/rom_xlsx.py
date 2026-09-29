# -*- coding: utf-8 -*-
"""Build <语言>构词分析.xlsx from rom_morph.js output, in the layout of the German sheet「1」."""
import io, os, sys, json, subprocess, openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from de_root_groups import cluster
lang = sys.argv[1]
d = json.loads(subprocess.check_output(["node", os.path.join(HERE, "rom_morph.js"), lang]).decode("utf-8"))
out = os.path.join(HERE, f"{d['name']}构词分析.xlsx")

gold = PatternFill("solid", fgColor="E8B64C"); alt = PatternFill("solid", fgColor="F6EFE3")
thin = Side(style="thin", color="B8A07A"); thick = Side(style="medium", color="8A6D33")
def header(ws, H, W):
    ws.append(H)
    for c in ws[1]: c.font = Font(bold=True); c.fill = gold
    for i, w in enumerate(W): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"

def family_sheet(ws, rows, with_group=True):
    fams = {}
    for r in rows: fams.setdefault(r["famKey"], {"key": r["famKey"], "zh": r["rootZh"], "en": r["rootEn"]})
    groups = cluster(list(fams.values())) if with_group else [[k] for k in fams]
    gid, glabel, order = {}, {}, []
    for g in groups: g.sort(key=lambda k: k.lower().lstrip("-("))
    for n, g in enumerate(sorted(groups, key=lambda g: g[0].lower().lstrip("-("))):
        for k in g: gid[k] = n
        glabel[n] = " · ".join(g) if len(g) > 1 else ""
        order += g
    pos = {k: i for i, k in enumerate(order)}
    rows = sorted(rows, key=lambda r: (pos[r["famKey"]], r["word"].lower()))
    means = {}
    for r in rows:
        m = means.setdefault(r["famKey"], ([], []))
        if r["rootZh"] and r["rootZh"] not in m[0]: m[0].append(r["rootZh"])
        if r["rootEn"] and r["rootEn"] not in m[1]: m[1].append(r["rootEn"])
    H = ["前缀", "复合前件", "连接成分", "核心词形", "复合后件", "余部", "后缀", "完整单词", "冠词", "中文", "English", "词根", "词根义中文", "词根义英语", "形近义近的词根组"]
    header(ws, H, [10, 11, 9, 11, 11, 6, 16, 20, 6, 20, 22, 14, 16, 18, 24])
    fam, g0, gcur, f0, shade = None, 2, None, 2, False
    def close_fam(end):
        if fam is not None and end > f0:
            for col in (12, 13, 14): ws.merge_cells(start_row=f0, start_column=col, end_row=end, end_column=col)
    def close_grp(end):
        if gcur is not None and glabel[gcur] and end > g0: ws.merge_cells(start_row=g0, start_column=15, end_row=end, end_column=15)
    for i, r in enumerate(rows):
        row = i + 2
        ws.append([r["pre"], r["comp1"], r["link"], r["core"], r["comp2"], r["tail"], r["suf"], r["word"], r["art"],
                   r["zh"], r["en"], r["famKey"], " / ".join(means[r["famKey"]][0]), " / ".join(means[r["famKey"]][1]), glabel[gid[r["famKey"]]]])
        if gid[r["famKey"]] != gcur:
            close_fam(row - 1); close_grp(row - 1)
            gcur, g0, fam, f0, shade = gid[r["famKey"]], row, r["famKey"], row, not shade
            for c in ws[row]: c.border = Border(top=thick)
        elif r["famKey"] != fam:
            close_fam(row - 1); fam, f0, shade = r["famKey"], row, not shade
            for c in ws[row]: c.border = Border(top=thin)
        if shade:
            for c in ws[row]: c.fill = alt
    close_fam(len(rows) + 1); close_grp(len(rows) + 1)
    for row in ws.iter_rows(min_row=2, min_col=12, max_col=15):
        for c in row: c.alignment = Alignment(vertical="center", wrap_text=True)
    return sum(1 for v in glabel.values() if v), len(fams)

wb = openpyxl.Workbook(); wb.remove(wb.active)
ng, nf = family_sheet(wb.create_sheet("1"), d["rows"])
family_sheet(wb.create_sheet("② 无缀词(光词根或只有词尾)"), d["plain"], with_group=False)

ws = wb.create_sheet("③ 词根总表")
header(ws, ["词根", "词根义中文", "词根义英语", "词数", "出现过的前缀", "出现过的后缀", "例词", "有词源故事"], [16, 18, 20, 7, 30, 34, 50, 8])
fam = {}
for r in d["rows"] + d["plain"]:
    f = fam.setdefault(r["famKey"], {"zh": r["rootZh"], "en": r["rootEn"], "n": 0, "pre": [], "suf": [], "ex": [], "st": r["hasStory"]})
    f["n"] += 1
    for x in filter(None, r["pre"].split(" + ")):
        if x not in f["pre"]: f["pre"].append(x)
    for x in filter(None, r["suf"].split(" + ")):
        if x not in f["suf"]: f["suf"].append(x)
    if len(f["ex"]) < 6: f["ex"].append(r["word"])
for k in sorted(fam, key=lambda k: k.lower().lstrip("-(")):
    f = fam[k]; ws.append([k, f["zh"], f["en"], f["n"], " ".join(f["pre"]), " ".join(f["suf"]), " · ".join(f["ex"]), f["st"]])

ws = wb.create_sheet("④ 前缀表"); header(ws, ["前缀", "意思", "本表词数", "表内例词"], [10, 30, 9, 60])
for r in sorted(d["preTable"], key=lambda r: -r[2]): ws.append(r)
ws = wb.create_sheet("⑤ 后缀表"); header(ws, ["后缀", "意思", "类型", "本表词数", "表内例词"], [11, 30, 12, 9, 60])
for r in sorted(d["sufTable"], key=lambda r: -r[3]): ws.append(r)
ws = wb.create_sheet("⑥ 词组"); header(ws, ["词条", "中文", "English", "词根ID", "词根"], [26, 22, 24, 12, 16])
for r in d["phrases"]: ws.append(r)
ws = wb.create_sheet("⑦ 待归词根"); header(ws, ["单词", "中文", "English", "原分类ID", "原分类"], [22, 20, 24, 12, 20])
for r in d["noRoot"]: ws.append(r)
ws = wb.create_sheet("⑧ 未切分"); header(ws, ["词根ID", "词根", "单词", "中文"], [12, 16, 22, 24])
for r in d["skipped"]: ws.append(r)
s = d["stat"]
ws = wb.create_sheet("⑨ 怎么读这张表"); header(ws, ["栏目", "说明"], [18, 96])
for r in [
    ["表「1」", "带前缀、派生后缀或复合成分的词。同一词根家族连成一块,词根和词根义各合并一格;形近又义近的家族排在一起,最后一列合并写出这一组。"],
    ["② 无缀词", "没有前缀、没有派生后缀的词。只带屈折词尾(-o/-a/-us/-er 不定式等)的也放这里,词尾仍写在「后缀」列。"],
    ["前缀", "包括同化形式:拉丁 ad- 在不同辅音前变成 ac-/af-/ag-/al-/an-/ap-/ar-/as-/at-,con- 变成 com-/col-/cor-/co-;意思一栏写明它是谁的变体。"],
    ["后缀", "「类型」区分:名词 / 形容词 / 动词 / 副词 = 派生后缀;构词中缀 = 词根和后缀之间的分词词干或中缀(accus-at-io 的 at);词尾 = 屈折词尾。"],
    ["复合前件 / 复合后件", "另一个独立的词根,多为希腊语复合(bio-logie、tele-phone)或拉丁复合(agri-cultura)。"],
    ["余部", "自动切分后剩下的 1–2 个字母,多为连接元音或词干变体,留着方便核对。"],
    ["词根来源", "词库已标词根的直接用;原来只按主题收录的词按最长词根自动匹配,例如在 ③ 中可见,数量见下。"],
    ["统计", f"共 {s['words']} 词:带缀 {s['affixed']},无缀 {s['plain']};{nf} 个词根家族,{ng} 个形近义近组;换元音识别 {s['ablaut']};自动匹配词根 {s['auto']};词组 {s['phrases']};待归词根 {s['noRoot']};未切分 {s['skipped']}。"],
]: ws.append(r)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(out)
print(f"{d['name']}: {out}  带缀 {s['affixed']} / 无缀 {s['plain']} / 家族 {nf} / 形近义近组 {ng}")
